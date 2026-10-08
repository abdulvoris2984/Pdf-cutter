import sys
import zipfile
from pathlib import Path

import pytest
from pypdf import PdfReader, PdfWriter
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import cutter  # noqa: E402
from cutter import CutterError  # noqa: E402


# --------------------------------------------------------------- planning --


def test_plan_by_size_even():
    assert cutter.plan_by_size(80, 20) == [(1, 20), (21, 40), (41, 60), (61, 80)]


def test_plan_by_size_uneven_last_part_is_shorter():
    assert cutter.plan_by_size(85, 20) == [(1, 20), (21, 40), (41, 60), (61, 80), (81, 85)]


def test_plan_by_size_bigger_than_document():
    assert cutter.plan_by_size(5, 20) == [(1, 5)]


def test_plan_by_count_even():
    assert cutter.plan_by_count(80, 4) == [(1, 20), (21, 40), (41, 60), (61, 80)]


def test_plan_by_count_spreads_remainder_over_first_parts():
    assert cutter.plan_by_count(10, 3) == [(1, 4), (5, 7), (8, 10)]


def test_plan_by_count_too_many_parts():
    with pytest.raises(CutterError):
        cutter.plan_by_count(3, 4)


@pytest.mark.parametrize(
    "text, expected",
    [
        ("1-20, 21-40", [(1, 20), (21, 40)]),
        ("1-25;26-50; 51-", [(1, 25), (26, 50), (51, 80)]),
        ("5", [(5, 5)]),
        ("1 – 10, 11—12", [(1, 10), (11, 12)]),
        ("1-80, 1-10", [(1, 80), (1, 10)]),
    ],
)
def test_parse_ranges(text, expected):
    assert cutter.parse_ranges(text, 80) == expected


@pytest.mark.parametrize("text", ["", "abc", "0-5", "1-81", "90-", "20-10", "1-2-3"])
def test_parse_ranges_rejects_bad_input(text):
    with pytest.raises(CutterError):
        cutter.parse_ranges(text, 80)


def test_errors_use_the_right_word():
    with pytest.raises(CutterError, match="outside the document .it has slides 1-12"):
        cutter.make_plan(12, "ranges", "90-", unit="slides")
    with pytest.raises(CutterError, match="only 12 slides"):
        cutter.make_plan(12, "count", "20", unit="slides")


def test_make_plan_dispatch_and_errors():
    assert cutter.make_plan(80, "size", "20") == cutter.plan_by_size(80, 20)
    assert cutter.make_plan(80, "count", " 4 ") == cutter.plan_by_count(80, 4)
    assert cutter.make_plan(80, "ranges", "1-5") == [(1, 5)]
    for bad in ["", "x", "0", "-3"]:
        with pytest.raises(CutterError):
            cutter.make_plan(80, "size", bad)


def test_output_paths(tmp_path):
    src = tmp_path / "exam.pdf"
    assert [p.name for p in cutter.output_paths(src, 4)] == [
        "exam1.pdf", "exam2.pdf", "exam3.pdf", "exam4.pdf",
    ]
    assert cutter.output_paths(src, 1)[0].parent == tmp_path
    other = tmp_path / "out"
    assert cutter.output_paths(src, 2, other)[1] == other / "exam2.pdf"


def test_unsupported_files(tmp_path):
    for name in ["notes.ppt", "notes.docx"]:
        path = tmp_path / name
        path.write_bytes(b"x")
        with pytest.raises(CutterError):
            cutter.count_pages(path)


# -------------------------------------------------------------------- PDF --


def make_pdf(path: Path, pages: int) -> Path:
    writer = PdfWriter()
    for i in range(pages):
        # Different widths let us check which original page ended up where.
        writer.add_blank_page(width=100 + i, height=200)
    with open(path, "wb") as fh:
        writer.write(fh)
    return path


def widths(path: Path) -> list[int]:
    return [int(page.mediabox.width) for page in PdfReader(path).pages]


def test_split_pdf_80_into_4(tmp_path):
    src = make_pdf(tmp_path / "exam.pdf", 80)
    assert cutter.count_pages(src) == 80

    seen = []
    outputs = cutter.split_file(src, cutter.plan_by_size(80, 20), progress=lambda d, t: seen.append((d, t)))

    assert [p.name for p in outputs] == ["exam1.pdf", "exam2.pdf", "exam3.pdf", "exam4.pdf"]
    for i, out in enumerate(outputs):
        assert widths(out) == list(range(100 + 20 * i, 120 + 20 * i))
    assert seen == [(1, 4), (2, 4), (3, 4), (4, 4)]


def test_split_pdf_into_other_folder(tmp_path):
    src = make_pdf(tmp_path / "exam.pdf", 5)
    outputs = cutter.split_file(src, [(2, 3), (5, 5)], out_dir=tmp_path / "new" / "folder")
    assert widths(outputs[0]) == [101, 102]
    assert widths(outputs[1]) == [104]
    assert outputs[0].parent == tmp_path / "new" / "folder"


def test_split_pdf_without_user_password(tmp_path):
    plain = make_pdf(tmp_path / "plain.pdf", 4)
    writer = PdfWriter(clone_from=plain)
    writer.encrypt(user_password="", owner_password="owner", algorithm="AES-128")
    locked = tmp_path / "locked.pdf"
    with open(locked, "wb") as fh:
        writer.write(fh)

    assert cutter.count_pages(locked) == 4
    outputs = cutter.split_file(locked, [(1, 2), (3, 4)])
    assert widths(outputs[1]) == [102, 103]


def test_password_protected_pdf_gives_friendly_error(tmp_path):
    writer = PdfWriter(clone_from=make_pdf(tmp_path / "plain.pdf", 2))
    writer.encrypt(user_password="secret", algorithm="AES-128")
    locked = tmp_path / "locked.pdf"
    with open(locked, "wb") as fh:
        writer.write(fh)
    with pytest.raises(CutterError, match="password"):
        cutter.count_pages(locked)


def test_broken_pdf_gives_friendly_error(tmp_path):
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"this is not a pdf")
    with pytest.raises(CutterError):
        cutter.count_pages(broken)


# ------------------------------------------------------------------- PPTX --

P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def make_pptx(path: Path, slides: int, with_sections: bool = False) -> Path:
    prs = Presentation()
    for i in range(1, slides + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        slide.shapes.title.text = f"Slide {i}"
        slide.notes_slide.notes_text_frame.text = f"Notes {i}"
    if with_sections:
        _add_sections(prs, [list(range(0, slides // 2)), list(range(slides // 2, slides))])
    prs.save(path)
    return path


def _add_sections(prs, groups):
    """Add PowerPoint 2010 sections, the way PowerPoint stores them."""
    from lxml import etree

    ids = [sld_id.get("id") for sld_id in prs.slides._sldIdLst.sldId_lst]
    pres = prs.part._element
    p = pres.nsmap.get("p", "http://schemas.openxmlformats.org/presentationml/2006/main")
    ext_lst = pres.find(f"{{{p}}}extLst")
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, f"{{{p}}}extLst")
    ext = etree.SubElement(ext_lst, f"{{{p}}}ext", uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    section_lst = etree.SubElement(ext, f"{{{P14_NS}}}sectionLst", nsmap={"p14": P14_NS})
    for n, group in enumerate(groups):
        section = etree.SubElement(section_lst, f"{{{P14_NS}}}section", name=f"Part {n}", id=f"{{0000000{n}-0000-0000-0000-000000000000}}")
        lst = etree.SubElement(section, f"{{{P14_NS}}}sldIdLst")
        for index in group:
            etree.SubElement(lst, f"{{{P14_NS}}}sldId", id=ids[index])


def titles(path: Path) -> list[str]:
    return [slide.shapes.title.text for slide in Presentation(path).slides]


def test_split_pptx(tmp_path):
    src = make_pptx(tmp_path / "lecture.pptx", 10)
    assert cutter.count_pages(src) == 10

    outputs = cutter.split_file(src, cutter.plan_by_count(10, 3))

    assert [p.name for p in outputs] == ["lecture1.pptx", "lecture2.pptx", "lecture3.pptx"]
    assert titles(outputs[0]) == [f"Slide {i}" for i in range(1, 5)]
    assert titles(outputs[1]) == [f"Slide {i}" for i in range(5, 8)]
    assert titles(outputs[2]) == [f"Slide {i}" for i in range(8, 11)]
    # Speaker notes travel with their slides.
    notes = [s.notes_slide.notes_text_frame.text for s in Presentation(outputs[2]).slides]
    assert notes == ["Notes 8", "Notes 9", "Notes 10"]


def test_split_pptx_drops_removed_slides_from_the_file(tmp_path):
    src = make_pptx(tmp_path / "lecture.pptx", 6)
    out = cutter.split_file(src, [(3, 4)])[0]
    with zipfile.ZipFile(out) as zf:
        slide_parts = [n for n in zf.namelist() if n.startswith("ppt/slides/slide")]
        notes_parts = [n for n in zf.namelist() if n.startswith("ppt/notesSlides/notesSlide")]
    assert len(slide_parts) == 2
    assert len(notes_parts) == 2


def test_split_pptx_cleans_up_sections(tmp_path):
    src = make_pptx(tmp_path / "lecture.pptx", 6, with_sections=True)
    out = cutter.split_file(src, [(2, 3)])[0]

    prs = Presentation(out)
    kept_ids = {sld_id.get("id") for sld_id in prs.slides._sldIdLst.sldId_lst}
    section_ids = {el.get("id") for el in prs.part._element.iter(f"{{{P14_NS}}}sldId")}
    assert titles(out) == ["Slide 2", "Slide 3"]
    assert section_ids == kept_ids


def test_self_test_passes():
    cutter.self_test()
