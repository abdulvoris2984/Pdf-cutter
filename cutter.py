"""Splitting logic for PDF Cutter (no GUI code in here, so it is easy to test).

Page/slide numbers are 1-based and ranges are inclusive, the same way you would
say them out loud: "pages 1-20" means pages 1 through 20.
"""

from __future__ import annotations

from pathlib import Path
from typing import Callable, Iterable

SUPPORTED_EXTENSIONS = (".pdf", ".pptx")

Range = tuple[int, int]
ProgressCallback = Callable[[int, int], None]  # (parts_done, parts_total)


class CutterError(Exception):
    """An error with a message that is safe to show to the user."""


# --------------------------------------------------------------------------- #
# Planning: turn the user's choice into a list of page ranges
# --------------------------------------------------------------------------- #


def plan_by_size(total: int, pages_per_part: int) -> list[Range]:
    """Chunks of `pages_per_part` pages; the last chunk may be shorter."""
    if pages_per_part < 1:
        raise CutterError("Pages per file must be at least 1.")
    return [
        (start, min(start + pages_per_part - 1, total))
        for start in range(1, total + 1, pages_per_part)
    ]


def plan_by_count(total: int, parts: int, unit: str = "pages") -> list[Range]:
    """Split into `parts` files of (almost) equal size.

    When the pages don't divide evenly, the first files get one extra page,
    e.g. 10 pages into 3 files -> 4, 3, 3.
    """
    if parts < 1:
        raise CutterError("Number of files must be at least 1.")
    if parts > total:
        raise CutterError(f"Can't make {parts} files out of only {total} {unit}.")
    base, extra = divmod(total, parts)
    ranges, start = [], 1
    for i in range(parts):
        size = base + (1 if i < extra else 0)
        ranges.append((start, start + size - 1))
        start += size
    return ranges


def parse_ranges(text: str, total: int, unit: str = "pages") -> list[Range]:
    """Parse custom ranges like "1-20, 21-35, 36-80" (one output file per item).

    A single number ("5") is a one-page file, and an open end ("61-") runs to the
    last page. Semicolons work as separators too.
    """
    items = [item.strip() for item in text.replace(";", ",").split(",")]
    items = [item for item in items if item]
    if not items:
        raise CutterError("Type at least one range, e.g. 1-20, 21-40.")

    ranges = []
    for item in items:
        normalized = item.replace("–", "-").replace("—", "-").replace(" ", "")
        try:
            if "-" in normalized:
                start_text, end_text = normalized.split("-", 1)
                start = int(start_text)
                end = int(end_text) if end_text else total
            else:
                start = end = int(normalized)
        except ValueError:
            raise CutterError(f'"{item}" is not a valid range. Use something like 1-20.') from None
        if start < 1 or start > total or end > total:
            raise CutterError(f'"{item}" is outside the document (it has {unit} 1-{total}).')
        if start > end:
            raise CutterError(f'"{item}" goes backwards. Write the smaller number first.')
        ranges.append((start, end))
    return ranges


def make_plan(total: int, mode: str, value: str, unit: str = "pages") -> list[Range]:
    """Dispatch on the GUI's mode: "size", "count" or "ranges".

    `unit` ("pages" or "slides") is only used in error messages.
    """
    if total < 1:
        raise CutterError(f"This file has no {unit}.")
    if mode == "ranges":
        return parse_ranges(value, total, unit)
    try:
        number = int(str(value).strip())
    except ValueError:
        raise CutterError("Please type a whole number.") from None
    if mode == "size":
        return plan_by_size(total, number)
    if mode == "count":
        return plan_by_count(total, number, unit)
    raise ValueError(f"Unknown mode: {mode}")


def output_paths(source: Path, parts: int, out_dir: Path | None = None) -> list[Path]:
    """exam.pdf -> exam1.pdf, exam2.pdf, ... (next to the source unless out_dir is given).

    Names that already end in a number get an underscore, so lecture03.pptx
    becomes lecture03_1.pptx instead of the confusing lecture031.pptx.
    """
    source = Path(source)
    folder = Path(out_dir) if out_dir else source.parent
    stem = source.stem + ("_" if source.stem[-1:].isdigit() else "")
    return [folder / f"{stem}{i}{source.suffix}" for i in range(1, parts + 1)]


# --------------------------------------------------------------------------- #
# Reading and writing files
# --------------------------------------------------------------------------- #


def check_supported(path: Path) -> None:
    suffix = Path(path).suffix.lower()
    if suffix == ".ppt":
        raise CutterError("Old .ppt files aren't supported. Open it in PowerPoint and save it as .pptx first.")
    if suffix not in SUPPORTED_EXTENSIONS:
        raise CutterError(f"{Path(path).name}: only PDF and PPTX files are supported.")


def count_pages(path: Path) -> int:
    """Number of pages (PDF) or slides (PPTX)."""
    path = Path(path)
    check_supported(path)
    try:
        if path.suffix.lower() == ".pdf":
            return len(_open_pdf(path).pages)
        return len(_open_pptx(path).slides)
    except CutterError:
        raise
    except Exception as exc:
        raise CutterError(f"Couldn't read {path.name}: {exc}") from exc


def split_file(
    source: Path,
    ranges: list[Range],
    out_dir: Path | None = None,
    progress: ProgressCallback | None = None,
) -> list[Path]:
    """Write one file per range and return the paths that were written."""
    source = Path(source)
    check_supported(source)
    if not ranges:
        raise CutterError("Nothing to cut: no page ranges were given.")
    targets = output_paths(source, len(ranges), out_dir)
    targets[0].parent.mkdir(parents=True, exist_ok=True)
    try:
        if source.suffix.lower() == ".pdf":
            _split_pdf(source, ranges, targets, progress)
        else:
            _split_pptx(source, ranges, targets, progress)
    except CutterError:
        raise
    except Exception as exc:
        raise CutterError(f"Couldn't cut {source.name}: {exc}") from exc
    return targets


# --- PDF -------------------------------------------------------------------- #


def _open_pdf(path: Path):
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        # Many lecture PDFs are "encrypted" only to block copying and open
        # without a password. Those work; password-protected ones don't.
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            raise CutterError(f"{path.name} is password-protected. Remove the password first.")
    return reader


def _split_pdf(source: Path, ranges: list[Range], targets: list[Path], progress) -> None:
    from pypdf import PdfWriter

    reader = _open_pdf(source)
    for done, ((start, end), target) in enumerate(zip(ranges, targets), start=1):
        writer = PdfWriter()
        # append() also carries over bookmarks that point into these pages.
        writer.append(reader, pages=(start - 1, end))
        with open(target, "wb") as fh:
            writer.write(fh)
        if progress:
            progress(done, len(ranges))


# --- PPTX ------------------------------------------------------------------- #

_P14_SLD_ID = "{http://schemas.microsoft.com/office/powerpoint/2010/main}sldId"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def _open_pptx(path: Path):
    from pptx import Presentation

    return Presentation(path)


def _keep_only_slides(prs, keep: Iterable[int]) -> None:
    """Delete every slide whose 0-based index is not in `keep`.

    python-pptx has no public "delete slide" API, so this unlinks the slides from
    presentation.xml. Unlinked slides (and images only they use) are not written
    when the file is saved.
    """
    keep = set(keep)
    pres_xml = prs.part._element
    sld_id_lst = prs.slides._sldIdLst
    removed_ids = set()

    for index, sld_id in enumerate(list(sld_id_lst.sldId_lst)):
        if index in keep:
            continue
        rId = sld_id.rId
        removed_ids.add(sld_id.get("id"))
        sld_id_lst.remove(sld_id)
        # Drop any other reference to the slide (e.g. custom slide shows) so the
        # relationship really goes away.
        for element in [el for el in pres_xml.iter() if el.get(_R_ID) == rId]:
            element.getparent().remove(element)
        prs.part.drop_rel(rId)

    # Sections (PowerPoint 2010+) list slides by id; stale entries make
    # PowerPoint offer to "repair" the file.
    for element in [el for el in pres_xml.iter(_P14_SLD_ID) if el.get("id") in removed_ids]:
        element.getparent().remove(element)


def _split_pptx(source: Path, ranges: list[Range], targets: list[Path], progress) -> None:
    for done, ((start, end), target) in enumerate(zip(ranges, targets), start=1):
        prs = _open_pptx(source)  # fresh copy per part
        _keep_only_slides(prs, range(start - 1, end))
        prs.save(target)
        if progress:
            progress(done, len(ranges))


def self_test() -> None:
    """End-to-end check used by the build to prove the packaged app really works.

    Raises if anything is broken (for example a library missing from the .exe).
    """
    import tempfile

    from pptx import Presentation
    from pypdf import PdfWriter

    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        writer = PdfWriter()
        for _ in range(10):
            writer.add_blank_page(width=200, height=200)
        writer.encrypt(user_password="", owner_password="owner", algorithm="AES-128")
        with open(folder / "test.pdf", "wb") as fh:
            writer.write(fh)
        prs = Presentation()
        for _ in range(10):
            prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(folder / "test.pptx")

        for name in ("test.pdf", "test.pptx"):
            source = folder / name
            outputs = split_file(source, plan_by_size(count_pages(source), 4))
            counts = [count_pages(path) for path in outputs]
            if counts != [4, 4, 2]:
                raise RuntimeError(f"{name}: expected parts of 4, 4, 2 but got {counts}")
