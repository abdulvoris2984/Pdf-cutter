"""Regenerate the images in docs/ (screenshots, demo GIF and social preview).

    python scripts/make_screenshots.py              # Windows / macOS
    xvfb-run -a python scripts/make_screenshots.py  # Linux without a desktop

It opens the real app, clicks through it and takes screenshots along the way.
"""

from __future__ import annotations

import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import customtkinter as ctk  # noqa: E402
from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageGrab  # noqa: E402
from pptx import Presentation  # noqa: E402
from pypdf import PdfWriter  # noqa: E402

import cutter  # noqa: E402
from ui.window import PdfCutterWindow  # noqa: E402

DOCS = ROOT / "docs"
FONT_DIRS = [Path("/usr/share/fonts/opentype/inter"), Path.home() / "Library/Fonts", Path("C:/Windows/Fonts")]


# ----------------------------------------------------------------- driving --


def make_samples(folder: Path) -> tuple[Path, Path]:
    pdf = folder / "Linear_Algebra_Exam.pdf"
    writer = PdfWriter()
    for _ in range(80):
        writer.add_blank_page(width=612, height=792)
    with open(pdf, "wb") as fh:
        writer.write(fh)
    pptx = folder / "Machine_Learning_Lecture_3.pptx"
    prs = Presentation()
    for _ in range(36):
        prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(pptx)
    return pdf, pptx


class Camera:
    def __init__(self, window: PdfCutterWindow):
        self.window = window
        self.frames: list[tuple[Image.Image, int]] = []

    def wait(self, seconds: float = 0.5) -> None:
        end = time.monotonic() + seconds
        while time.monotonic() < end:
            self.window.update()
            time.sleep(0.01)

    def grab(self) -> Image.Image:
        self.wait()
        w = self.window
        x, y = w.winfo_rootx(), w.winfo_rooty()
        return ImageGrab.grab(bbox=(x, y, x + w.winfo_width(), y + w.winfo_height())).convert("RGB")

    def frame(self, ms: int) -> None:
        self.frames.append((self.grab(), ms))


def set_mode(window: PdfCutterWindow, label: str) -> None:
    window.mode_switch.set(label)
    window._set_mode(label)


def slowed(split_file):
    """Make cutting take a moment so the progress bar shows up in the GIF."""

    def wrapper(source, ranges, out_dir=None, progress=None):
        def report(done, total):
            time.sleep(0.4)
            if progress:
                progress(done, total)

        return split_file(source, ranges, out_dir, report)

    return wrapper


def record() -> tuple[Image.Image, Image.Image, list[tuple[Image.Image, int]]]:
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp) / "Exams"
        folder.mkdir()
        pdf, pptx = make_samples(folder)
        ctk.set_appearance_mode("Light")
        window = PdfCutterWindow()
        window.geometry("1040x680+0+0")
        cam = Camera(window)

        # Still screenshots.
        window.add_files([pdf, pptx])
        light = cam.grab()
        ctk.set_appearance_mode("Dark")
        dark = cam.grab()
        ctk.set_appearance_mode("Light")
        window.clear_files()

        # Demo animation.
        cam.frame(1400)
        window.drop_zone.set_highlight(True)
        cam.frame(700)
        window.drop_zone.set_highlight(False)
        window.add_files([pdf, pptx])
        cam.frame(2200)
        set_mode(window, "Number of files")
        window.values["count"].set("3")
        cam.frame(1800)
        window.remove_file(pptx)
        cam.frame(1000)
        set_mode(window, "Custom ranges")
        for text in ["1-30", "1-30, 31-55", "1-30, 31-55, 56-"]:
            window.values["ranges"].set(text)
            cam.frame(600)
        cam.frame(1000)
        cutter.split_file = slowed(cutter.split_file)
        window.start_cut()
        while window.busy:
            cam.frame(300)
        cam.frame(2400)
        ctk.set_appearance_mode("Dark")
        cam.frame(2600)
        window.destroy()
    return light, dark, cam.frames


# ---------------------------------------------------------------- styling --


def rounded_mask(size: tuple[int, int], radius: int, scale: int = 4) -> Image.Image:
    big = Image.new("L", (size[0] * scale, size[1] * scale), 0)
    ImageDraw.Draw(big).rounded_rectangle((0, 0, *big.size), radius=radius * scale, fill=255)
    return big.resize(size, Image.LANCZOS)


def framed(shot: Image.Image, radius: int = 14, margin: int = 48) -> Image.Image:
    """Rounded corners, a hairline border and a soft drop shadow."""
    w, h = shot.size
    canvas = Image.new("RGBA", (w + 2 * margin, h + 2 * margin), (0, 0, 0, 0))
    mask = rounded_mask((w, h), radius)
    shadow = Image.new("RGBA", canvas.size, (15, 23, 42, 0))
    shadow_alpha = Image.new("L", canvas.size, 0)
    shadow_alpha.paste(mask.point(lambda a: a * 0.32), (margin, margin + 12))
    shadow.putalpha(shadow_alpha.filter(ImageFilter.GaussianBlur(22)))
    canvas.alpha_composite(shadow)
    canvas.paste(shot, (margin, margin), mask)
    border = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    ImageDraw.Draw(border).rounded_rectangle(
        (margin, margin, margin + w - 1, margin + h - 1), radius=radius, outline=(120, 120, 140, 70), width=1
    )
    return Image.alpha_composite(canvas, border)


def font(weight: str, size: int) -> ImageFont.FreeTypeFont:
    for folder in FONT_DIRS:
        for name in (f"Inter-{weight}.otf", f"Inter-{weight}.ttf", "segoeuib.ttf" if weight == "Bold" else "segoeui.ttf"):
            if (folder / name).exists():
                return ImageFont.truetype(str(folder / name), size)
    return ImageFont.load_default(size)


def social_preview(light: Image.Image) -> Image.Image:
    """1280x640 image for GitHub's social preview (Settings → General)."""
    width, height = 1280, 640
    gradient = Image.new("RGB", (2, 2))
    gradient.putdata([(129, 140, 248), (99, 102, 241), (99, 102, 241), (55, 48, 163)])
    card = gradient.resize((width, height), Image.BILINEAR).convert("RGBA")

    icon = Image.open(ROOT / "assets" / "icon.png").resize((112, 112), Image.LANCZOS)
    card.alpha_composite(icon, (80, 150))
    draw = ImageDraw.Draw(card)
    draw.text((80, 290), "PDF Cutter", font=font("Bold", 68), fill="white")
    draw.multiline_text(
        (82, 382), "Split PDFs and slide decks\ninto smaller files.", font=font("Medium", 32),
        fill=(255, 255, 255, 225), spacing=10,
    )
    draw.text((82, 500), "Windows  ·  macOS  ·  Linux", font=font("SemiBold", 22), fill=(224, 231, 255, 230))

    shot = framed(light.resize((760, round(760 * light.height / light.width)), Image.LANCZOS), radius=12, margin=40)
    card.alpha_composite(shot, (width - shot.width + 150, (height - shot.height) // 2 + 10))
    return card.convert("RGB")


def save_gif(frames: list[tuple[Image.Image, int]], path: Path, width: int = 880) -> None:
    smalls = [image.resize((width, round(width * image.height / image.width)), Image.LANCZOS) for image, _ in frames]
    # One shared palette, built from all frames, keeps colors accurate and stops flicker.
    strip = Image.new("RGB", (width, sum(s.height for s in smalls)))
    y = 0
    for small in smalls:
        strip.paste(small, (0, y))
        y += small.height
    palette = strip.quantize(colors=256, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE)
    images = [small.quantize(palette=palette, dither=Image.Dither.NONE) for small in smalls]
    durations = [ms for _, ms in frames]
    images[0].save(path, save_all=True, append_images=images[1:], duration=durations, loop=0, optimize=True)


def main() -> None:
    DOCS.mkdir(exist_ok=True)
    light, dark, frames = record()
    framed(light).save(DOCS / "screenshot-light.png", optimize=True)
    framed(dark).save(DOCS / "screenshot-dark.png", optimize=True)
    social_preview(light).save(DOCS / "social-preview.png", optimize=True)
    save_gif(frames, DOCS / "demo.gif")
    for name in sorted(p.name for p in DOCS.iterdir()):
        print(f"docs/{name}: {(DOCS / name).stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
